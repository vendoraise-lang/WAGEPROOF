from pathlib import Path
from openpyxl import load_workbook
from decimal import Decimal, ROUND_HALF_UP
import subprocess,tempfile,json,zipfile
R=Path(__file__).resolve().parents[1]
def cents(x):return float(Decimal(str(x)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP))
cases=[('blank',{},{6:0,10:0,15:0,18:0,24:0,12:None,17:None}),('markup_vs_margin',{'Estimate':{'A4':'Labor','C4':10,'D4':100,'F4':.25}}, {4:1250,7:1000,11:250,12:.2}),('overhead_change_actuals',{'Estimate':{'A4':'Base','C4':1,'D4':8000,'F4':.25},'Project Setup':{'B7':1000,'B8':900,'B9':.1},'Change Orders':{'C4':'Approved','D4':2000,'E4':1200,'C5':'Pending','D5':9999,'E5':999,'C6':'Rejected','D6':8888},'Actual Costs':{'D4':1,'E4':8500,'D5':1,'E5':1250},'Payments':{'C4':6000,'C5':2000}},{4:10000,5:2000,6:12000,8:1200,10:10200,11:1800,12:.15,13:9750,15:10650,16:1350,18:8000,19:4000,20:-2650,21:-450,22:1020}),('deductive_change_refund',{'Estimate':{'C4':1,'D4':800,'F4':.25},'Change Orders':{'C4':'Approved','D4':-100,'E4':-50},'Payments':{'C4':1000,'C5':-100}},{6:900,10:750,11:150,18:900,19:0}),('zero_revenue_costs',{'Actual Costs':{'D4':1,'E4':200}},{6:0,15:200,16:-200,12:None,17:None}),('rounding',{'Estimate':{'C4':3,'D4':.335,'F4':.25},'Actual Costs':{'D4':3,'E4':.335}},{7:1.01,4:1.26,13:1.01}),('last_rows_overpayment',{'Estimate':{'C103':2,'D103':50,'F103':.2},'Actual Costs':{'D203':1,'E203':110},'Change Orders':{'C103':'Approved','D103':30,'E103':10},'Payments':{'C103':200}},{6:150,15:110,18:200,19:-50}),('incomplete_entries',{'Estimate':{'A4':'Missing quantity','D4':100},'Actual Costs':{'B4':'Missing cost','D4':1},'Change Orders':{'C4':'Approved','D4':500},'Payments':{'B4':'Missing payment'}},{24:4})]
with tempfile.TemporaryDirectory() as tmp:
 p=Path(tmp);out=p/'recalc';out.mkdir()
 for name,updates,expected in cases:
  w=load_workbook(R/'deliverables/Contractor_Profit_Toolkit.xlsx')
  for sh,vals in updates.items():
   for cell,val in vals.items():w[sh][cell]=val
  w.save(p/(name+'.xlsx'))
 for f in (R/'deliverables').glob('*.xlsx'):(p/f.name).write_bytes(f.read_bytes())
 run=subprocess.run(['soffice','-env:UserInstallation=file://'+tmp+'/profile','--headless','--convert-to','xlsx','--outdir',str(out),*[str(x) for x in p.glob('*.xlsx')]],capture_output=True,text=True,timeout=600)
 assert run.returncode==0,run.stderr
 results=[]
 for name,updates,expected in cases:
  w=load_workbook(out/(name+'.xlsx'),data_only=True)
  for row,val in expected.items():
   got=w['Dashboard'][f'B{row}'].value
   if val is None:assert got in (None,''),(name,row,got)
   else:assert isinstance(got,(int,float)) and abs(got-val)<1e-8,(name,row,got,val)
  for s in w:
   for row in s:
    for c in row:assert c.data_type!='e',(name,s.title,c.coordinate,c.value)
  results.append({'scenario':name,'result':'PASS','asserted_metrics':expected})
 for f in (R/'deliverables').glob('*.xlsx'):
  rec=out/f.name;w=load_workbook(rec,data_only=True)
  for s in w:
   for row in s:
    for c in row:assert c.data_type!='e',(f.name,s.title,c.coordinate,c.value)
  # Keep original styles/validation, transplant only actual recalculated cached values from Calc XML.
  import xml.etree.ElementTree as ET
  ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
  original=zipfile.ZipFile(f);calc=zipfile.ZipFile(rec);payload={n:original.read(n) for n in original.namelist()};original.close()
  for n in payload:
   if n.startswith('xl/worksheets/sheet') and n.endswith('.xml'):
    tree=ET.fromstring(payload[n]);ct=ET.fromstring(calc.read(n));cv={c.attrib['r']:c for c in ct.findall('.//s:sheetData/s:row/s:c',ns)}
    for c in tree.findall('.//s:sheetData/s:row/s:c',ns):
     if c.find('s:f',ns) is None:continue
     cc=cv[c.attrib['r']]
     for v in list(c):
      if v.tag.endswith('}v'):c.remove(v)
     v=cc.find('s:v',ns)
     if v is not None:c.append(v)
     if 't' in cc.attrib:c.set('t',cc.attrib['t'])
     elif 't' in c.attrib:del c.attrib['t']
    payload[n]=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
  calc.close()
  with zipfile.ZipFile(f,'w',zipfile.ZIP_DEFLATED) as z:
   for n,data in payload.items():z.writestr(n,data)
  with zipfile.ZipFile(f) as z:
   assert not any('vbaProject' in n or 'externalLinks/' in n for n in z.namelist())
  w=load_workbook(f);assert not w._external_links
  assert sum(c.data_type=='f' for s in w for row in s for c in row)==521
 (R/'qa/test_results.json').write_text(json.dumps({'engine':subprocess.check_output(['soffice','--version'],text=True).strip(),'scenarios':results,'formula_count_per_workbook':521,'errors':0,'macro_external_link_scan':'PASS','note':'Recalculated in LibreOffice. Microsoft Excel UI not available; compatibility not independently verified in Excel.'},indent=2))
 print('PASS: 8 recalculated scenarios, both final files, 521 formulas/file, no formula errors, macros or external links.')
