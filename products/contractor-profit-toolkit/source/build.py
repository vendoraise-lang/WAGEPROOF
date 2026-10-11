from pathlib import Path
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.workbook.properties import CalcProperties
import json
ROOT=Path(__file__).resolve().parents[1]
INPUTS={'Project Setup': ['B4:B13'], 'Estimate':['A4:D103','F4:F103','H4:H103'],'Actual Costs':['A4:E203','G4:H203'],'Change Orders':['A4:E103','G4:G103'],'Payments':['A4:E103']}
MONEY='"$"#,##0.00;[Red]("$"#,##0.00);"$"0.00'
def validate(s,typ,f1,f2,area,operator='between'):
 d=DataValidation(type=typ,operator=operator,formula1=f1,formula2=f2,allow_blank=True)
 d.showErrorMessage=True;d.errorTitle='Check this entry';d.error='Use the allowed value. Read the column notes.';d.errorStyle='stop';d.showInputMessage=True;d.promptTitle='Input cell';d.prompt='Enter numbers without currency symbols. Clear unused rows.';s.add_data_validation(d);d.add(area)
def build(demo):
 w=load_workbook(ROOT/'source/original.xlsx');w.calculation=CalcProperties(calcId=0,fullCalcOnLoad=True)
 for s in w:
  s.sheet_view.showGridLines=False;s.freeze_panes='C4' if s.max_column>3 else 'B4'
  s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_LETTER;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
  s.print_title_rows='1:3';s.print_options.horizontalCentered=True
  s.sheet_properties.tabColor='16324F';s.row_dimensions[1].height=34;s.row_dimensions[3].height=32
  s.data_validations.dataValidation=[]
  for row in s:
   for c in row:
    c.font=Font(name='Calibri',size=11,color='243746');c.alignment=Alignment(vertical='center',wrap_text=True);c.protection=Protection(locked=True)
    if c.row==1: c.fill=PatternFill('solid',fgColor='16324F');c.font=Font(name='Calibri',size=16,bold=True,color='FFFFFF')
    elif c.row==3: c.fill=PatternFill('solid',fgColor='DCE7ED');c.font=Font(name='Calibri',size=10,bold=True,color='16324F')
    elif c.row%2==0:c.fill=PatternFill('solid',fgColor='F4F7F9')
    if c.data_type=='f':c.font=Font(name='Calibri',size=11,bold=True,color='16324F')
  for area in INPUTS.get(s.title,[]):
   for row in s[area]:
    for c in row:c.protection=Protection(locked=False);c.fill=PatternFill('solid',fgColor='E8F2FF');c.font=Font(name='Calibri',size=11,color='1457A0')
  s.protection.sheet=True;s.protection.selectLockedCells=False;s.protection.selectUnlockedCells=False;s.protection.autoFilter=False
  if s.max_column>3:
   s.auto_filter.ref=f'A3:{chr(64+s.max_column)}{s.max_row}'
   for col in 'ABCDEFGH':s.column_dimensions[col].width=22
   s.column_dimensions['A'].width=29;s.column_dimensions['B'].width=30;s.column_dimensions[chr(64+s.max_column)].width=45
  else:
   s.column_dimensions['A'].width=36;s.column_dimensions['B'].width=31;s.column_dimensions['C'].width=76
  for r in range(4,s.max_row+1):s.row_dimensions[r].height=60 if s.title=='START HERE' else 48 if s.title=='Project Setup' else 32
 for r in range(4,104):
  w['Estimate'][f'E{r}']=f'=IF(OR(C{r}="",D{r}=""),"",ROUND(C{r}*D{r},2))'
  w['Estimate'][f'G{r}']=f'=IF(E{r}="","",ROUND(E{r}*(1+IF(F{r}="",0,F{r})),2))'
  w['Change Orders'][f'F{r}']=f'=IF(C{r}="Approved",IF(OR(D{r}="",E{r}=""),"",ROUND(D{r}-E{r},2)),0)'
 for r in range(4,204):w['Actual Costs'][f'F{r}']=f'=IF(OR(D{r}="",E{r}=""),"",ROUND(D{r}*E{r},2))'
 s=w['Dashboard'];s['A11']='Budgeted job contribution';s['A16']='Contract less costs incurred';s['A20']='Receipts less incurred costs';s['C20']='NOT cash flow: expenses may be unpaid. No expense-payment ledger is included.'
 s['B12']='=IF(B6=0,"",B11/B6)';s['B17']='=IF(B6=0,"",B16/B6)';s['B22']="=ROUND(B10*'Project Setup'!B9,2)"
 s['C11']='Contract revenue less estimated direct costs and attributable overhead; excludes corporate overhead and taxes.'
 s['C12']='Contribution / contracted revenue. Blank when revenue is zero; markup is cost-based.'
 s['C17']='Provisional contribution / contract. Blank when revenue is zero; no completion forecast.'
 s['C19']='Negative means receipts exceed contract. Not accounts receivable or an invoice aging report.'
 s['A24']='Entry warnings';s['B24']='=SUMPRODUCT(--(Estimate!A4:A103<>""),--(Estimate!C4:C103=""))+SUMPRODUCT(--(Estimate!A4:A103<>""),--(Estimate!D4:D103=""))+SUMPRODUCT(--(\'Actual Costs\'!B4:B203<>""),--(\'Actual Costs\'!D4:D203=""))+SUMPRODUCT(--(\'Actual Costs\'!B4:B203<>""),--(\'Actual Costs\'!E4:E203=""))+SUMPRODUCT(--(\'Change Orders\'!C4:C103="Approved"),--(\'Change Orders\'!D4:D103=""))+SUMPRODUCT(--(\'Change Orders\'!C4:C103="Approved"),--(\'Change Orders\'!E4:E103=""))+SUMPRODUCT(--(Payments!B4:B103<>""),--(Payments!C4:C103=""))';s['C24']='Missing required numeric fields on described entries. Counts fields, not rows. Review before relying on totals.'
 s['B28']='Provisional until every cost is recorded. Blank margin means zero contract revenue.'
 s.conditional_formatting.add('B11:B11',CellIsRule(operator='lessThan',formula=['0'],fill=PatternFill('solid',fgColor='FADBD8')))
 s.conditional_formatting.add('B24',CellIsRule(operator='greaterThan',formula=['0'],fill=PatternFill('solid',fgColor='FADBD8')))
 s.conditional_formatting.add('B12',CellIsRule(operator='lessThan',formula=['B23'],fill=PatternFill('solid',fgColor='FFF0C2')))
 for r in range(4,24):s[f'B{r}'].number_format='0.0%' if r in [12,17,23] else MONEY
 s['B24'].number_format='0'
 setup=w['Project Setup'];setup['A13']='Target contribution margin';setup['C13']='Benchmark after project overhead; excludes unallocated company overhead and tax.';setup['C10']='Excluded informational allowance. Enter all other figures consistently net of customer sales tax.'
 for col in ['B7','B8','B10']:setup[col].number_format=MONEY
 for col in ['B9','B13']:setup[col].number_format='0.0%'
 for area in ['B7:B8','B10']:validate(setup,'decimal','0','1000000000',area)
 validate(setup,'decimal','0','1','B9');validate(setup,'decimal','0','0.9999','B13')
 for name,cols in [('Estimate',['D','E','G']),('Actual Costs',['E','F']),('Change Orders',['D','E','F']),('Payments',['C'])]:
  sh=w[name]
  for col in cols:
   for r in range(4,sh.max_row+1):sh[f'{col}{r}'].number_format=MONEY
 validate(w['Estimate'],'decimal','0','1000000000','C4:D103');validate(w['Estimate'],'decimal','0','10','F4:F103')
 for r in range(4,104):w['Estimate'][f'F{r}'].number_format='0.0%'
 validate(w['Actual Costs'],'decimal','0','1000000000','D4:E203')
 validate(w['Change Orders'],'decimal','-1000000000','1000000000','D4:E103')
 validate(w['Payments'],'decimal','-1000000000','1000000000','C4:C103')
 for name,area,items in [('Estimate','B4:B103','Materials,Labor,Equipment,Subcontract,Other'),('Actual Costs','C4:C203','Materials,Labor,Equipment,Subcontract,Other'),('Actual Costs','H4:H203','Base,Change'),('Change Orders','C4:C103','Pending,Approved,Rejected'),('Payments','D4:D103','Bank transfer,Check,Card,Cash,Other')]:validate(w[name],'list','"'+items+'"',None,area)
 for name in ['Actual Costs','Payments']:
  validate(w[name],'date','DATE(2000,1,1)','DATE(2100,12,31)',f'A4:A{w[name].max_row}')
  for r in range(4,w[name].max_row+1):w[name][f'A{r}'].number_format='mmm d, yyyy'
 validate(setup,'date','DATE(2000,1,1)','DATE(2100,12,31)','B12')
 w['START HERE']['A1']='CONTRACTOR PROFIT TOOLKIT | 1.2';w['START HERE']['B9']='Review Dashboard warnings and contribution margins.';w['START HERE']['C9']='Incomplete job costs overstate profit. Receipts less incurred costs is NOT cash flow.'
 w['START HERE']['C10']='Customer tax excluded. Costs include nonrecoverable purchase taxes. Negative payment = refund.'
 w['START HERE']['C11']='Fixed capacity: 100 estimate, 200 costs, 100 changes, 100 payment rows. Do not add rows.'
 w['START HERE']['C13']='Edit blue cells. Sheet protection prevents accidents, not unauthorized access. No password needed to unprotect.'
 w['RELEASE NOTES']['B2']='1.2 production package';w['RELEASE NOTES']['B7']='One purchasing business may reuse for its own jobs. No redistribution or resale. See included license.'
 if not demo:
  for name,areas in INPUTS.items():
   for area in areas:
    for row in w[name][area]:
     for c in row:c.value=None
  for cell,val in {'B7':0,'B8':0,'B9':0,'B10':0,'B13':0.20}.items():setup[cell]=val
 else:setup['B8']=500
 # Reapply formula styling after replacement.
 for sh in w:
  for row in sh:
   for c in row:
    if c.data_type=='f':c.font=Font(name='Calibri',size=11,bold=True,color='16324F');c.protection=Protection(locked=True)
  sh.print_area=f'A1:{chr(64+sh.max_column)}{sh.max_row}'
 w.active=w.sheetnames.index('START HERE');w.save(ROOT/'deliverables'/('Contractor_Profit_Toolkit_Example.xlsx' if demo else 'Contractor_Profit_Toolkit.xlsx'))
if __name__=='__main__':
 build(False);build(True)
 w=load_workbook(ROOT/'deliverables/Contractor_Profit_Toolkit.xlsx')
 inventory=[{'sheet':s.title,'cell':c.coordinate,'formula':c.value} for s in w for row in s for c in row if c.data_type=='f']
 (ROOT/'qa/formula_inventory.json').write_text(json.dumps(inventory,indent=2))
 print('Built blank and fictional example;',len(inventory),'formulas inventoried.')
